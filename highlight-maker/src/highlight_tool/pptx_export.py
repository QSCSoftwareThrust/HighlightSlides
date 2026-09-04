from __future__ import annotations

import os
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt


FIELD_NAMES = [
    "Highlight Title",
    "Scientific Achievement",
    "Significance and Impact",
    "Research Details",
    "Recommended Figure",
    "Figure Caption",
    "Citation",
    "DOI",
    "Acknowledgements",
    "User Facilities",
    "Institution Logos",
    "NQISRC Intellectual Role",
    "NQISRC Funding Role",
    "Funding Institution Contributions",
]

SLIDE1_ACHIEVEMENT = (5484242, 846572, 6536773, 1385422)
SLIDE1_IMPACT = (5484242, 2136544, 6536773, 1385422)
SLIDE1_DETAILS = (5428304, 3279528, 6536773, 1385422)
SLIDE1_FIGURE = (548770, 1107088, 4935472, 2850432)
SLIDE1_CAPTION = (633893, 3957520, 4709288, 1077218)
SLIDE1_REFERENCE = (408791, 5065329, 5075451, 646331)
SLIDE2_ROLE = (577016, 1114915, 10516969, 1077218)
SLIDE2_CONTRIBUTIONS = (576498, 2219044, 11149337, 1600000)


def parse_key_information(path: Path) -> dict[str, str]:
    text = path.read_text(encoding="utf-8")
    pattern = re.compile(r"^##\s+(.+?)\s*$", re.MULTILINE)
    matches = list(pattern.finditer(text))
    fields: dict[str, str] = {}

    for index, match in enumerate(matches):
        name = match.group(1).strip()
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        if name in FIELD_NAMES:
            fields[name] = text[start:end].strip()

    return {name: fields.get(name, "").strip() for name in FIELD_NAMES}


def _set_text(shape, text: str, *, font_size: int = 15, bold_first_line: bool = True) -> None:
    shape.text_frame.clear()
    lines = text.splitlines() or [""]

    first = shape.text_frame.paragraphs[0]
    first.text = lines[0]
    first.font.size = Pt(font_size)
    first.font.color.rgb = RGBColor(0, 0, 0)
    first.font.bold = bold_first_line

    for line in lines[1:]:
        paragraph = shape.text_frame.add_paragraph()
        clean = line.strip()
        is_bullet = clean.startswith("- ")
        paragraph.text = clean[2:] if is_bullet else clean
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = RGBColor(0, 0, 0)
        paragraph.level = 0
        if is_bullet:
            paragraph._p.get_or_add_pPr().set("marL", "228600")
            paragraph._p.get_or_add_pPr().set("indent", "-114300")

    shape.text_frame.word_wrap = True
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = RGBColor(0, 0, 0)


def _set_title(shape, title: str) -> None:
    shape.text_frame.clear()
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = title or "Untitled Highlight"
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(0, 0, 0)
    paragraph.alignment = PP_ALIGN.LEFT


def _delete_shape(shape) -> None:
    shape.element.getparent().remove(shape.element)


def _add_text_box(slide, left, top, width, height, text: str, *, font_size: int = 16):
    box = slide.shapes.add_textbox(left, top, width, height)
    _set_text(box, text, font_size=font_size)
    return box


def _delete_slides_after(prs: Presentation, keep_count: int) -> None:
    slide_ids = list(prs.slides._sldIdLst)
    for slide_id in slide_ids[keep_count:]:
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)


def _image_candidates(text: str) -> list[str]:
    candidates: list[str] = []
    candidates.extend(re.findall(r"!\[[^\]]*\]\(([^)]+?\.(?:png|jpg|jpeg))\)", text, flags=re.IGNORECASE))
    candidates.extend(re.findall(r"\[[^\]]+\]\(([^)]+?\.(?:png|jpg|jpeg))\)", text, flags=re.IGNORECASE))
    candidates.extend(re.findall(r"`([^`]+?\.(?:png|jpg|jpeg))`", text, flags=re.IGNORECASE))
    candidates.extend(re.findall(r"([A-Za-z0-9_./-]+\.(?:png|jpg|jpeg))", text, flags=re.IGNORECASE))
    return candidates


def _resolve_project_image(project_dir: Path, candidate: str) -> Path | None:
    path = Path(candidate.strip())
    if path.is_absolute() and path.exists():
        return path
    candidates = [path, *_white_variant_candidates(path)]
    roots = [
        project_dir,
        project_dir / "logos",
        project_dir / "extracted" / "figures",
        Path.cwd(),
        Path.cwd() / "assets" / "logos",
        Path.cwd().parent / "assets" / "logos",
        Path("/work/highlight-slides/assets/logos"),
        Path(__file__).resolve().parents[3] / "assets" / "logos",
        Path(__file__).resolve().parents[4] / "assets" / "logos",
    ]
    for root in roots:
        for candidate_path in candidates:
            resolved = root / candidate_path
            if resolved.exists():
                return resolved
    matches = sorted(project_dir.rglob(path.name))
    return matches[0] if matches else None


def _white_variant_candidates(path: Path) -> list[Path]:
    if path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
        return []
    stem = path.stem
    if stem.endswith("-white"):
        return []
    return [path.with_name(f"{stem}-white.png")]


def _first_existing_figure(project_dir: Path, recommended: str) -> Path | None:
    candidates = _image_candidates(recommended)

    for candidate in candidates:
        page_match = re.search(r"page-(\d{3})\.(?:png|jpg|jpeg)$", candidate, flags=re.IGNORECASE)
        if page_match:
            crop_matches = sorted(
                (project_dir / "extracted" / "figures" / "crops").glob(f"page-{page_match.group(1)}-figure-*.png")
            )
            if crop_matches:
                return crop_matches[0]

    for candidate in candidates:
        path = Path(candidate)
        if path.is_absolute() and path.exists():
            return path
        resolved = project_dir / path
        if resolved.exists():
            return resolved
        figure_resolved = project_dir / "extracted" / "figures" / path
        if figure_resolved.exists():
            return figure_resolved

    for candidate in candidates:
        basename = Path(candidate).name
        matches = sorted((project_dir / "extracted" / "figures").rglob(basename))
        if matches:
            return matches[0]

    return None


def _add_contained_picture(slide, image_path: Path, left, top, width, height) -> None:
    with Image.open(image_path) as image:
        image_width, image_height = image.size

    image_ratio = image_width / image_height
    box_ratio = width / height

    if image_ratio > box_ratio:
        final_width = width
        final_height = int(width / image_ratio)
    else:
        final_height = height
        final_width = int(height * image_ratio)

    final_left = left + int((width - final_width) / 2)
    final_top = top + int((height - final_height) / 2)
    slide.shapes.add_picture(str(image_path), final_left, final_top, final_width, final_height)


def _logo_paths(project_dir: Path, logo_text: str) -> list[Path]:
    paths: list[Path] = []
    seen: set[Path] = set()

    for candidate in _image_candidates(logo_text):
        resolved = _resolve_project_image(project_dir, candidate)
        if resolved and resolved not in seen:
            paths.append(resolved)
            seen.add(resolved)
    return paths


def _required_qsc_logo(project_dir: Path) -> Path:
    for candidate in _qsc_logo_candidates(project_dir):
        if candidate.exists():
            return candidate
    searched = "\n".join(f"- {path}" for path in _qsc_logo_candidates(project_dir))
    raise FileNotFoundError(
        "Missing required QSC logo. Add the official logo as "
        "`HighlightSlides/assets/qsc-logo.png` or set HIGHLIGHT_QSC_LOGO.\n"
        f"Searched:\n{searched}"
    )


def _qsc_logo_candidates(project_dir: Path) -> list[Path]:
    env_path = os.environ.get("HIGHLIGHT_QSC_LOGO")
    candidates = []
    if env_path:
        candidates.append(Path(env_path))
    candidates.extend(
        [
            project_dir / "logos" / "qsc-logo.png",
            project_dir / "logos" / "qsc-logo-white.png",
            project_dir / "logos" / "qsc-logo.jpg",
            Path.cwd() / "assets" / "qsc-logo.png",
            Path.cwd() / "assets" / "qsc-logo-white.png",
            Path.cwd() / "assets" / "qsc-logo.jpg",
            Path.cwd().parent / "assets" / "qsc-logo.png",
            Path.cwd().parent / "assets" / "qsc-logo-white.png",
            Path.cwd().parent / "assets" / "qsc-logo.jpg",
            Path("/work/highlight-slides/assets/qsc-logo.png"),
            Path("/work/highlight-slides/assets/qsc-logo-white.png"),
            Path(__file__).resolve().parents[3] / "assets" / "qsc-logo.png",
            Path(__file__).resolve().parents[3] / "assets" / "qsc-logo-white.png",
            Path(__file__).resolve().parents[4] / "assets" / "qsc-logo.png",
            Path(__file__).resolve().parents[4] / "assets" / "qsc-logo-white.png",
        ]
    )
    return candidates


def _add_logo_row(slide, logo_paths: list[Path], left, top, width, height) -> None:
    max_logos = min(len(logo_paths), 6)
    if max_logos == 0:
        return

    logos = logo_paths[:max_logos]
    gap = int(width * 0.025)
    slot_width = int((width - gap * (max_logos - 1)) / max_logos)
    max_height = int(height * 0.72)
    min_height = int(height * 0.35)

    for index, logo_path in enumerate(logos):
        with Image.open(logo_path) as image:
            image_width, image_height = image.size
        if image_width <= 0 or image_height <= 0:
            continue

        ratio = image_width / image_height
        final_height = max_height
        final_width = int(final_height * ratio)
        if final_width > slot_width:
            final_width = slot_width
            final_height = int(final_width / ratio)
        if final_height < min_height and final_height > 0 and final_width > 0:
            scale = min(min_height / final_height, slot_width / final_width)
            final_height = int(final_height * scale)
            final_width = int(final_width * scale)

        slot_left = left + index * (slot_width + gap)
        final_left = slot_left + int((slot_width - final_width) / 2)
        final_top = top + int((height - final_height) / 2)
        slide.shapes.add_picture(str(logo_path), final_left, final_top, final_width, final_height)


def _add_qsc_and_partner_logos(slide, qsc_logo: Path, partner_logos: list[Path], left, top, width, height) -> None:
    qsc_slot_width = int(width * 0.24)
    gap = int(width * 0.04)
    partner_width = max(0, width - qsc_slot_width - gap)

    _add_contained_picture(slide, qsc_logo, left, top, qsc_slot_width, height)

    if partner_logos and partner_width > 0:
        partner_count = min(len(partner_logos), 5)
        partner_area_width = min(partner_width, partner_count * qsc_slot_width + max(0, partner_count - 1) * gap)
        _add_logo_row(slide, partner_logos, left + qsc_slot_width + gap, top, partner_area_width, height)


def _text_with_heading(heading: str, body: str) -> str:
    body = body.strip() or "Unknown"
    return f"{heading}\n{body}"


def _bulleted_text(heading: str, body: str) -> str:
    lines = [line.rstrip() for line in body.splitlines() if line.strip()]
    if not lines:
        return f"{heading}\n- Unknown"
    normalized = []
    for line in lines:
        stripped = line.strip()
        normalized.append(stripped if stripped.startswith("- ") else f"- {stripped}")
    return "\n".join([heading, *normalized])


def _shorten_citation(citation: str) -> str:
    citation = " ".join(citation.split())
    if not citation:
        return "Unknown"

    title_match = re.search(r'["“](.+?)["”]', citation)
    title = title_match.group(1).strip().rstrip(",") if title_match else ""
    before_title = citation[: title_match.start()].strip(" ,") if title_match else ""
    after_title = citation[title_match.end() :].strip(" ,") if title_match else ""

    author_count = 0
    if before_title:
        author_count = before_title.count(",") + before_title.count(" and ") + 1

    if title and before_title and (author_count > 6 or len(citation) > 260):
        first_author = before_title.split(",")[0].strip()
        suffix = f", {after_title}" if after_title else ""
        return f'{first_author} et al., "{title}"{suffix}'

    return citation


def _reference_text(project_dir: Path, citation: str, doi: str) -> str:
    reference = _shorten_citation(citation)
    doi = doi.strip()
    if doi and doi.lower() != "unknown" and doi not in reference:
        reference = f"{reference} DOI: {doi}."
    arxiv_id = _arxiv_id(project_dir)
    if arxiv_id and not re.search(r"arxiv\s*:\s*\d{4}\.\d+", reference, flags=re.IGNORECASE):
        reference = re.sub(r"\barXiv\b(?!\s*:)", f"arXiv:{arxiv_id}", reference, flags=re.IGNORECASE)
        if f"arXiv:{arxiv_id}" not in reference:
            reference = f"{reference} arXiv:{arxiv_id}."
    return "\n".join(["Reference(s)", reference])


def _arxiv_id(project_dir: Path) -> str:
    source_path = project_dir / "source" / "arxiv_source.txt"
    if source_path.exists():
        for line in source_path.read_text(encoding="utf-8", errors="replace").splitlines():
            if line.startswith("id:"):
                return line.split(":", 1)[1].strip()
    match = re.search(r"arxiv-(\d{4})-(\d+)", project_dir.name)
    if match:
        return f"{match.group(1)}.{match.group(2)}"
    return ""


def _role_explanation(role: str) -> str:
    normalized = role.strip().lower()
    if normalized == "lead":
        return "QSC/NQISRC was the lead institution contributing to the published result."
    if normalized == "collaborator":
        return "QSC/NQISRC was one of a group of collaborators with roughly equal contributions."
    if normalized == "minor":
        return "QSC/NQISRC contributed to work led by a person or project outside the center."
    return "Unknown from the paper; do not infer a role without support."


def _funding_explanation(role: str) -> str:
    normalized = role.strip().lower()
    if normalized == "sole":
        return "QSC/NQISRC was the sole funding source for this work."
    if normalized == "majority":
        return "Other sources contributed, but QSC/NQISRC made up most of the funding."
    if normalized == "minority":
        return "QSC/NQISRC funding was a minority of the overall funding."
    return "Unknown from the paper; do not infer a funding share without support."


def _slide2_role_text(fields: dict[str, str]) -> str:
    intellectual_role = fields["NQISRC Intellectual Role"] or "Unknown"
    funding_role = fields["NQISRC Funding Role"] or "Unknown"
    return "\n".join(
        [
            "NQISRC Intellectual Role:",
            intellectual_role,
            "",
            f"NQISRC Funding Role: {funding_role}",
        ]
    )


def _slide2_contribution_text(fields: dict[str, str]) -> str:
    return _bulleted_text(
        "Explanation of Funding Institution Contributions:",
        fields["Funding Institution Contributions"],
    )


def build_highlight_deck(
    project_dir: Path,
    *,
    template_path: Path,
    output_path: Path | None = None,
) -> Path:
    key_info_path = project_dir / "paper_key_information.md"
    if not key_info_path.exists():
        raise FileNotFoundError(f"Missing {key_info_path}")
    if not template_path.exists():
        raise FileNotFoundError(f"Missing template PowerPoint: {template_path}")

    fields = parse_key_information(key_info_path)
    prs = Presentation(str(template_path))
    _delete_slides_after(prs, 2)

    slide1 = prs.slides[0]
    _set_title(slide1.shapes[0], fields["Highlight Title"])
    for shape, geometry in [
        (slide1.shapes[1], SLIDE1_ACHIEVEMENT),
        (slide1.shapes[2], SLIDE1_IMPACT),
        (slide1.shapes[3], SLIDE1_DETAILS),
        (slide1.shapes[4], SLIDE1_CAPTION),
        (slide1.shapes[5], SLIDE1_REFERENCE),
    ]:
        shape.left, shape.top, shape.width, shape.height = geometry

    _set_text(slide1.shapes[1], _text_with_heading("Scientific Achievement", fields["Scientific Achievement"]), font_size=16)
    _set_text(slide1.shapes[2], _text_with_heading("Significance and Impact", fields["Significance and Impact"]), font_size=16)
    _set_text(slide1.shapes[3], _bulleted_text("Research Details", fields["Research Details"]), font_size=16)
    _set_text(slide1.shapes[4], fields["Figure Caption"] or "Short figure caption", font_size=16, bold_first_line=False)
    _set_text(slide1.shapes[5], _reference_text(project_dir, fields["Citation"], fields["DOI"]), font_size=12)

    image_box = slide1.shapes[7]
    image_path = _first_existing_figure(project_dir, fields["Recommended Figure"])
    if image_path:
        _delete_shape(image_box)
        left, top, width, height = SLIDE1_FIGURE
        _add_contained_picture(slide1, image_path, left, top, width, height)
    else:
        image_box.left, image_box.top, image_box.width, image_box.height = SLIDE1_FIGURE
        _set_text(image_box, _text_with_heading("Image(s)", fields["Recommended Figure"]), font_size=14)

    logo_box = slide1.shapes[6]
    qsc_logo = _required_qsc_logo(project_dir)
    partner_logos = _logo_paths(project_dir, fields["Institution Logos"])
    left, top, width, height = logo_box.left, logo_box.top, logo_box.width, logo_box.height
    _delete_shape(logo_box)
    _add_qsc_and_partner_logos(slide1, qsc_logo, partner_logos, left, top, width, height)

    slide2 = prs.slides[1]
    _set_title(slide2.shapes[0], f"{fields['Highlight Title'] or 'Highlight'}: NQISRC Roles")
    content_box = slide2.shapes[1]
    _delete_shape(content_box)

    _add_text_box(
        slide2,
        *SLIDE2_ROLE,
        _slide2_role_text(fields),
        font_size=24,
    )
    _add_text_box(
        slide2,
        *SLIDE2_CONTRIBUTIONS,
        _slide2_contribution_text(fields),
        font_size=20,
    )

    output = output_path or project_dir / "output" / "highlight_slides.pptx"
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output
